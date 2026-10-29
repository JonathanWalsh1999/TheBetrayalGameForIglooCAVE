"""Build The Betrayal: inject the sample questions into the template and write index.html + questions.xlsx.
Run:  python3 src/build.py
The questions.xlsx format is the same as the Monopoly-style game, so one question bank works in both."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)
SETTINGS = {"name": "The Betrayal", "pot": 500, "discuss": 3, "qper": 3, "secs": 45}

data = json.load(open(os.path.join(HERE, "sample_questions.json")))
data["settings"] = SETTINGS
tpl = open(os.path.join(HERE, "index.template.html"), encoding="utf-8").read()
assert "/*__SAMPLE_DATA__*/null" in tpl
open(os.path.join(OUT, "index.html"), "w", encoding="utf-8").write(
    tpl.replace("/*__SAMPLE_DATA__*/null", json.dumps(data, ensure_ascii=False)))

wb = Workbook()
help_ = wb.active; help_.title = "How to use"
rows = [
    ["The Betrayal – question bank"],
    [""],
    ["Questions sheet: one row per question."],
    ["Topic: optional (any code is fine – kept so this file also works in the Monopoly-style game)."],
    ["Level: 1 = easier, 2 = medium, 3 = harder. Round 1 uses level 1, later rounds use harder ones when there are enough."],
    ["Option A–D: leave blank for an open question (the tutor reveals the answer)."],
    ["Answer: the letter (A–D) for multiple choice, or the answer text for an open question."],
    ["Explanation: shown after the answer – keep it to one sentence."],
    [""],
    ["Settings sheet: change the numbers in the Value column."],
    ["How many questions? Each mission asks 3 (every house answers each). 4 rounds = 12 questions, so 20+ means no repeats; 40+ is ideal."],
]
for r in rows: help_.append(r)
help_["A1"].font = Font(bold=True, size=14); help_.column_dimensions["A"].width = 120

ws = wb.create_sheet("Questions")
head = ["Topic", "Level", "Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Explanation"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="5B1A24")
for q in data["questions"]:
    o = {x["L"]: x["text"] for x in q["opts"]}
    ws.append([q["topic"], q["level"], q["q"], o.get("A", ""), o.get("B", ""), o.get("C", ""), o.get("D", ""), q["ans"], q["exp"]])
for col, w in zip("ABCDEFGHI", [8, 7, 60, 24, 24, 24, 24, 9, 60]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

st = wb.create_sheet("Settings")
st.append(["Setting", "Value", "What it does"])
st.append(["Game name", SETTINGS["name"], "Shown on the walls"])
st.append(["Pot per correct answer", SETTINGS["pot"], "Pounds added to the prize pot for each house that answers a mission question right"])
st.append(["Questions per mission", SETTINGS["qper"], "How many questions each mission has (every house answers each one)"])
st.append(["Seconds per question", SETTINGS["secs"], "Thinking time before houses lock in their answer"])
st.append(["Discussion minutes", SETTINGS["discuss"], "Length of each Round Table discussion"])
for c in st[1]: c.font = Font(bold=True)
st.column_dimensions["A"].width = 26; st.column_dimensions["B"].width = 16; st.column_dimensions["C"].width = 70
wb.save(os.path.join(OUT, "questions.xlsx"))
print("questions:", len(data["questions"]))
